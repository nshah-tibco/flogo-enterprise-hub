# Build a PowerPoint (.pptx) customer-workshop deck for PSA (Port of Singapore Authority):
# "Flogo for Event Streaming at the Port" — positioning Flogo as a lightweight event-handling
# agent (Kafka / JMS) that COMPLEMENTS TIBCO Streaming, illustrated by a Reefer Cold-Chain
# Monitoring use case.
#
# Native shapes/tables => fully editable after importing into Google Slides.
#
# IMPORTANT (Google Slides compatibility): python-pptx with zero or float EMU extents produces
# files that Google Slides rejects as corrupt. Every shape/textbox/table below coerces all EMU
# offsets/extents to int(...) and nudges any zero extent to a small positive int.
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn

# ---- palette ----
INK   = RGBColor(0x0F,0x1E,0x2E)
MUTED = RGBColor(0x5B,0x6B,0x7B)
BG    = RGBColor(0xEE,0xF2,0xF6)
CARD  = RGBColor(0xFF,0xFF,0xFF)
LINE  = RGBColor(0xD8,0xE0,0xE8)
NAVY  = RGBColor(0x0B,0x20,0x38)
NAVY2 = RGBColor(0x12,0x3A,0x5E)
WHITE = RGBColor(0xFF,0xFF,0xFF)
ACCENT= RGBColor(0x0E,0xA5,0xA4)   # teal — Flogo / enrich
CARDDK= RGBColor(0x0E,0x2C,0x46)
LIGHTTEAL = RGBColor(0xEA,0xF6,0xF6)

PRODUCER = RGBColor(0x16,0xA3,0x4A)   # green — Flogo producer
KAFKA    = RGBColor(0x1F,0x29,0x37)   # near-black — Apache Kafka
CONSUMER = RGBColor(0x63,0x66,0xF1)   # indigo — Flogo consumer
DB       = RGBColor(0x64,0x74,0x8B)   # slate — PostgreSQL
ALERT    = RGBColor(0xDC,0x2A,0x2A)   # red — breach / detect
AMBER    = RGBColor(0xF5,0x9E,0x0B)   # amber — route / alerts
STREAM   = RGBColor(0xA8,0x55,0xF7)   # purple — TIBCO Streaming

FONT = "Segoe UI"
EMU = 914400
SW, SH = Inches(13.333), Inches(7.5)

prs = Presentation()
prs.slide_width = SW
prs.slide_height = SH
BLANK = prs.slide_layouts[6]

# ---- int-coercion helpers (Google Slides rejects float / zero EMU extents) ----
def _o(v):            # offset -> int EMU
    return int(round(float(v)))
def _e(v):            # extent -> int EMU, never zero
    return max(int(round(float(v))), 1)

def slide(bg=BG):
    s = prs.slides.add_slide(BLANK)
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0,0, _e(SW), _e(SH))
    r.fill.solid(); r.fill.fore_color.rgb = bg; r.line.fill.background()
    r.shadow.inherit = False
    return s

def _set_text(tf, text, size, color, bold=False, align=PP_ALIGN.LEFT, font=FONT, space_after=6):
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_after = Pt(space_after)
    runs = text if isinstance(text, list) else [(text, {})]
    for i,(t,opt) in enumerate(runs):
        r = p.add_run(); r.text = t
        f = r.font; f.size = Pt(opt.get("size",size)); f.name = font
        f.bold = opt.get("bold",bold); f.color.rgb = opt.get("color",color)
    return p

def textbox(s, x,y,w,h, text, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT,
            anchor=MSO_ANCHOR.TOP, font=FONT, space_after=6):
    tb = s.shapes.add_textbox(_o(x),_o(y),_e(w),_e(h)); tf = tb.text_frame
    tf.vertical_anchor = anchor
    tf.margin_left=0; tf.margin_right=0; tf.margin_top=0; tf.margin_bottom=0
    _set_text(tf, text, size, color, bold, align, font, space_after)
    return tb

def bullets(s, x,y,w,h, items, size=17, color=INK, gap=10, dot=ACCENT, font=FONT):
    tb = s.shapes.add_textbox(_o(x),_o(y),_e(w),_e(h)); tf = tb.text_frame; tf.word_wrap=True
    tf.margin_left=0; tf.margin_right=0; tf.margin_top=0; tf.margin_bottom=0
    first=True
    for it in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first=False
        p.space_after = Pt(gap)
        rb = p.add_run(); rb.text="■  "; rb.font.size=Pt(11); rb.font.color.rgb=dot; rb.font.name=font
        parts = it if isinstance(it,list) else [(it,{})]
        for t,opt in parts:
            r=p.add_run(); r.text=t; f=r.font
            f.size=Pt(opt.get("size",size)); f.name=font; f.bold=opt.get("bold",False)
            f.color.rgb=opt.get("color",color)
    return tb

def box(s, x,y,w,h, fill, radius=0.10, line=None, shadow=True):
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, _o(x),_o(y),_e(w),_e(h))
    try: shp.adjustments[0]=radius
    except Exception: pass
    shp.fill.solid(); shp.fill.fore_color.rgb=fill
    if line is None: shp.line.fill.background()
    else: shp.line.color.rgb=line; shp.line.width=Pt(1)
    shp.shadow.inherit=False
    if shadow:
        el = shp._element.spPr
        ef = el.makeelement(qn('a:effectLst'), {}); el.append(ef)
        sh = ef.makeelement(qn('a:outerShdw'), {'blurRad':'90000','dist':'40000','dir':'5400000','rotWithShape':'0'})
        ef.append(sh)
        clr = sh.makeelement(qn('a:srgbClr'), {'val':'0F1E2E'}); sh.append(clr)
        a = clr.makeelement(qn('a:alpha'), {'val':'18000'}); clr.append(a)
    return shp

def boxtext(s, shp, lines, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE):
    tf = shp.text_frame; tf.word_wrap=True; tf.vertical_anchor=anchor
    tf.margin_left=Pt(6); tf.margin_right=Pt(6); tf.margin_top=Pt(4); tf.margin_bottom=Pt(4)
    first=True
    for (t,size,color,bold) in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first=False; p.alignment=align
        r=p.add_run(); r.text=t; f=r.font
        f.size=Pt(size); f.color.rgb=color; f.bold=bold; f.name=FONT

def eyebrow(s, text, y=Inches(0.55), color=ACCENT):
    textbox(s, Inches(0.8), y, Inches(11.7), Inches(0.4),
            text.upper(), size=13, color=color, bold=True)

def title(s, text, y=Inches(0.95), color=INK, size=32):
    bar = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                             _o(Inches(0.8)), _o(y+Inches(0.12)), _e(Inches(0.5)), _e(Inches(0.10)))
    bar.fill.solid(); bar.fill.fore_color.rgb=ACCENT; bar.line.fill.background(); bar.shadow.inherit=False
    textbox(s, Inches(1.45), y, Inches(11.1), Inches(0.7), text, size=size, color=color, bold=True)

def footer(s, right="Reefer Cold-Chain · PSA Workshop", dark=False):
    c = RGBColor(0x9F,0xB6,0xC6) if dark else MUTED
    textbox(s, Inches(0.8), Inches(7.02), Inches(4), Inches(0.35),
            [("FLOGO", {"bold":True,"color":c}), (" · ", {"bold":True,"color":ACCENT}), ("EVENT STREAMING", {"bold":True,"color":c})],
            size=13)
    textbox(s, Inches(8.5), Inches(7.02), Inches(4.0), Inches(0.35), right, size=12, color=c, align=PP_ALIGN.RIGHT)

def chip(s, x, y, text, fill, w=None, h=Inches(0.42), fsize=13):
    w = w or Inches(0.02*len(text)+0.5)
    c = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, _o(x),_o(y),_e(w),_e(h))
    try: c.adjustments[0]=0.5
    except Exception: pass
    c.fill.solid(); c.fill.fore_color.rgb=fill; c.line.fill.background(); c.shadow.inherit=False
    boxtext(s, c, [(text,fsize,WHITE,True)])
    return c

def connector(s, x1,y1,x2,y2, color=RGBColor(0x41,0x58,0x6B), width=2.2, dashed=False, arrow=True):
    # OOXML requires integer EMU coords and NON-ZERO extents; Google Slides rejects zero/float.
    x1,y1,x2,y2 = _o(x1),_o(y1),_o(x2),_o(y2)
    if abs(x2-x1) < 9525: x2 = x1 + 9525   # ~0.01" — invisible, keeps cx>0
    if abs(y2-y1) < 9525: y2 = y1 + 9525   # keeps cy>0
    cn = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1,y1,x2,y2)
    cn.line.color.rgb=color; cn.line.width=Pt(width); cn.shadow.inherit=False
    ln = cn.line._get_or_add_ln()
    if dashed:
        d = ln.makeelement(qn('a:prstDash'), {'val':'dash'}); ln.append(d)
    if arrow:
        he = ln.makeelement(qn('a:tailEnd'), {'type':'triangle','w':'med','len':'med'}); ln.append(he)
    return cn

def card(s, x,y,w,h, topcolor=ACCENT):
    box(s, x,y,w,h, CARD, radius=0.06, line=LINE)
    top = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, _o(x),_o(y),_e(w),_e(Inches(0.07)))
    top.fill.solid(); top.fill.fore_color.rgb=topcolor; top.line.fill.background(); top.shadow.inherit=False

def table(s, x,y,w, rows, colw, header=True, fontsize=13.5, rowh=Inches(0.42)):
    nrows=len(rows); ncols=len(rows[0])
    gt = s.shapes.add_table(nrows,ncols,_o(x),_o(y),_e(w),_e(rowh*nrows)).table
    gt.first_row=header; gt.horz_banding=False
    for ci,cw in enumerate(colw): gt.columns[ci].width=_e(cw)
    for ri,row in enumerate(rows):
        gt.rows[ri].height=_e(rowh)
        for ci,val in enumerate(row):
            cell=gt.cell(ri,ci)
            cell.margin_left=Pt(8); cell.margin_right=Pt(6); cell.margin_top=Pt(3); cell.margin_bottom=Pt(3)
            cell.vertical_anchor=MSO_ANCHOR.MIDDLE
            if header and ri==0:
                cell.fill.solid(); cell.fill.fore_color.rgb=NAVY
                col=WHITE; bold=True; sz=12.5
            else:
                cell.fill.solid(); cell.fill.fore_color.rgb=CARD
                col=INK; bold=False; sz=fontsize
            tf=cell.text_frame; tf.word_wrap=True
            p=tf.paragraphs[0]; r=p.add_run(); r.text=str(val)
            f=r.font; f.size=Pt(sz); f.name=FONT; f.bold=bold; f.color.rgb=col
    return gt

# =========================================================================
# 1 TITLE
s = slide(NAVY)
band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0,0, _e(SW), _e(Inches(0.14)))
band.fill.solid(); band.fill.fore_color.rgb=ACCENT; band.line.fill.background(); band.shadow.inherit=False
textbox(s, Inches(0.9), Inches(1.35), Inches(11.5), Inches(0.5),
        "PSA WORKSHOP · TIBCO FLOGO ENTERPRISE", size=15, color=RGBColor(0x7F,0xD4,0xD1), bold=True)
textbox(s, Inches(0.9), Inches(1.95), Inches(11.8), Inches(2.0),
        "Flogo for Event Streaming\nat the Port", size=50, color=WHITE, bold=True)
textbox(s, Inches(0.9), Inches(4.15), Inches(11.2), Inches(1.2),
        "Lightweight event-handling agents for Kafka & JMS — complementing TIBCO Streaming.",
        size=21, color=RGBColor(0xC7,0xD7,0xE3))
for i,(t,c) in enumerate([("PSA Workshop",ACCENT),("Kafka",AMBER),("JMS / EMS",CONSUMER),("Event Handling",PRODUCER)]):
    chip(s, Inches(0.9+i*3.0), Inches(5.55), t, c, w=Inches(2.8))
footer(s, "Reefer Cold-Chain Monitoring — architecture workshop", dark=True)

# =========================================================================
# 2 THE OPPORTUNITY / CONTEXT
s = slide(); eyebrow(s,"The opportunity"); title(s,"Event handling at the world's busiest hub")
data=[("PSA scale",DB,"One of the world's largest transshipment hubs",
       "Thousands of containers move through PSA every day, throwing off a continuous, high-velocity stream of real-time operational events."),
      ("Since 2019",STREAM,"TIBCO Streaming already in production",
       "PSA has run TIBCO Streaming since 2019 for high-velocity, in-memory analytics and pattern detection on operational streams."),
      ("The ask",ACCENT,"A lightweight event-handling agent",
       "Ingest, enrich, route and bridge events across Kafka and JMS/EMS — close to where events are produced. Event-native, not API-centric.")]
cw=Inches(3.9); gap=Inches(0.25); x0=Inches(0.8); y0=Inches(2.0); h=Inches(3.7)
for i,(k,tc,hd,body) in enumerate(data):
    x=x0+i*(cw+gap); card(s,x,y0,cw,h,tc)
    textbox(s,x+Inches(0.28),y0+Inches(0.32),cw-Inches(0.5),Inches(0.35),k.upper(),size=12,color=MUTED,bold=True)
    textbox(s,x+Inches(0.28),y0+Inches(0.72),cw-Inches(0.5),Inches(1.0),hd,size=19,color=INK,bold=True)
    textbox(s,x+Inches(0.28),y0+Inches(1.85),cw-Inches(0.5),Inches(1.7),body,size=15,color=MUTED,space_after=0)
textbox(s,Inches(0.8),Inches(6.05),Inches(11.7),Inches(0.8),
        [("Position Flogo as the nimble edge agent: ",{"color":INK}),
         ("event-native ingestion, enrichment and routing that feeds — and acts on — TIBCO Streaming.",{"bold":True,"color":INK})],
        size=17)
footer(s)

# =========================================================================
# 3 COMPLEMENTARY, NOT COMPETING
s = slide(); eyebrow(s,"Complementary, not competing")
title(s,"Flogo ↔ TIBCO Streaming")
rows=[["TIBCO Streaming — since 2019","Flogo — new event agent"],
      ["High-throughput Complex Event Processing (CEP)","Lightweight edge ingestion of events"],
      ["Temporal / windowing operators & continuous queries","Message enrichment via DB / API lookups"],
      ["In-memory analytics on high-velocity streams","Content-based routing & filtering"],
      ["Statistical & ML operators","Protocol bridging: Kafka ⇄ JMS/EMS ⇄ REST"],
      ["Real-time pattern detection & alerting","Transformation & single-binary event agents"],
      ["Best for: real-time analytics & pattern detection","Best for: event handling, routing & integration glue"]]
gt = table(s, Inches(0.8), Inches(1.8), Inches(11.7), rows, [Inches(5.85),Inches(5.85)], rowh=Inches(0.53))
# emphasize the final "Best for" row
last = len(rows)-1
for ci in range(2):
    cell = gt.cell(last,ci); cell.fill.solid(); cell.fill.fore_color.rgb=LIGHTTEAL
    rr = cell.text_frame.paragraphs[0].runs[0]; rr.font.bold=True; rr.font.color.rgb=RGBColor(0x0B,0x6B,0x6A)
# Better-together banner
bb = box(s, Inches(0.8), Inches(5.85), Inches(11.7), Inches(1.0), NAVY, radius=0.06)
textbox(s, Inches(1.1), Inches(6.0), Inches(11.1), Inches(0.75),
        [("Better together:  ",{"bold":True,"color":RGBColor(0x7F,0xD4,0xD1)}),
         ("Flogo ingests, normalizes & routes events at the edge and feeds Streaming for heavy analytics — "
          "and Flogo also acts on Streaming's decisions: fire alerts, write to the database, notify downstream.",
          {"color":WHITE})], size=15)
footer(s)

# =========================================================================
# 4 USE CASE — REEFER COLD-CHAIN
s = slide(); eyebrow(s,"Use case"); title(s,"Reefer Cold-Chain Monitoring")
textbox(s, Inches(0.8), Inches(1.85), Inches(11.7), Inches(1.0),
        [("Refrigerated (“reefer”) containers carry temperature-sensitive cargo — pharma, food, perishables. ",{"color":INK}),
         ("A temperature excursion or power loss means spoilage, insurance claims and regulatory exposure.",{"bold":True,"color":INK})],
        size=18)
data=[("At stake",ALERT,"Spoiled cargo & claims",
       "Excursions or power loss ruin cargo, trigger insurance claims, and create regulatory exposure for the port and its customers."),
      ("The need",AMBER,"Monitor & alert instantly",
       "Continuously watch reefer telemetry and alert the moment a breach occurs — each container has a cargo-specific setpoint and tolerance."),
      ("Why streaming",ACCENT,"A natural event problem",
       "High volume, continuous telemetry that needs enrichment plus instant routing and alerting — exactly an event-streaming workload.")]
cw=Inches(3.9); gap=Inches(0.25); x0=Inches(0.8); y0=Inches(2.95); h=Inches(3.1)
for i,(k,tc,hd,body) in enumerate(data):
    x=x0+i*(cw+gap); card(s,x,y0,cw,h,tc)
    chip(s,x+Inches(0.28),y0+Inches(0.3),k,tc,w=Inches(1.9))
    textbox(s,x+Inches(0.28),y0+Inches(0.95),cw-Inches(0.5),Inches(0.6),hd,size=19,color=INK,bold=True)
    textbox(s,x+Inches(0.28),y0+Inches(1.6),cw-Inches(0.5),Inches(1.4),body,size=14.5,color=MUTED,space_after=0)
footer(s)

# =========================================================================
# 5 ARCHITECTURE (money slide)
s = slide(); eyebrow(s,"Architecture — Flogo on both sides of Kafka")
title(s,"Producer → Kafka → Consumer → alerts")

def cbox(x,y,w,h,fill,lines,radius=0.10):
    b=box(s,x,y,w,h,fill,radius=radius); boxtext(s,b,lines); return b

midy = Inches(3.15)
# Producer
cbox(Inches(0.55),Inches(2.4),Inches(2.5),Inches(1.55),PRODUCER,
     [("Flogo Producer",14.5,WHITE,True),("ReeferTelemetryPublisher",11,RGBColor(0xE6,0xF6,0xEC),False),
      ("timer-driven telemetry",10.5,RGBColor(0xE6,0xF6,0xEC),False),("simulator",10.5,RGBColor(0xE6,0xF6,0xEC),False)])
# Kafka (telemetry)
cbox(Inches(3.3),Inches(2.4),Inches(1.85),Inches(1.55),KAFKA,
     [("Apache Kafka",14,WHITE,True),("topic:",10.5,RGBColor(0xB8,0xC2,0xCC),False),
      ("reefer.telemetry",11.5,AMBER,True)])
# Consumer (big box with internal stages)
cbox(Inches(5.4),Inches(1.95),Inches(4.35),Inches(2.6),CONSUMER,
     [(" ",6,WHITE,False)])
textbox(s, Inches(5.6), Inches(2.08), Inches(4.0), Inches(0.6),
        [("Flogo Consumer",{"bold":True,"color":WHITE,"size":14.5})], size=14.5, align=PP_ALIGN.CENTER)
textbox(s, Inches(5.6), Inches(2.5), Inches(4.0), Inches(0.3),
        "ReeferMonitorProcessor", size=10.5, color=RGBColor(0xE3,0xE3,0xFF), align=PP_ALIGN.CENTER)
stg=[("Parse",DB),("Enrich",ACCENT),("Detect",ALERT),("Route",AMBER)]
sx=Inches(5.58); sw=Inches(0.95); sgap=Inches(0.08)
for i,(nm,cl) in enumerate(stg):
    chip(s, sx+i*(sw+sgap), Inches(2.95), nm, cl, w=sw, h=Inches(0.44), fsize=11.5)
textbox(s, Inches(5.6), Inches(3.55), Inches(4.0), Inches(0.35),
        "parse → enrich (DB) → detect breach → route", size=11, color=RGBColor(0xE3,0xE3,0xFF), align=PP_ALIGN.CENTER)
textbox(s, Inches(5.6), Inches(3.9), Inches(4.0), Inches(0.5),
        [("normal → log   ·   ",{"color":RGBColor(0xE3,0xE3,0xFF)}),("breach → alerts + audit",{"bold":True,"color":WHITE})],
        size=11, align=PP_ALIGN.CENTER)
# Outputs (stacked right)
ox=Inches(10.0); ow=Inches(2.8)
cbox(ox,Inches(1.95),ow,Inches(0.85),AMBER,
     [("Kafka · reefer.alerts",12.5,WHITE,True),("breach events published",10,RGBColor(0xFF,0xF3,0xDE),False)])
cbox(ox,Inches(2.98),ow,Inches(0.85),DB,
     [("Postgres · audit table",12.5,WHITE,True),("breach persisted",10,RGBColor(0xEC,0xF0,0xF4),False)])
cbox(ox,Inches(4.01),ow,Inches(0.55),RGBColor(0x41,0x58,0x6B),
     [("Log / console",11.5,WHITE,True)])
# Postgres enrichment source (below consumer)
cbox(Inches(5.4),Inches(5.05),Inches(4.35),Inches(0.95),DB,
     [("PostgreSQL — reefer_config",13,WHITE,True),
      ("setpoint · tolerance · cargo type · customer · destination",10.5,RGBColor(0xEC,0xF0,0xF4),False)])
# connectors
connector(s, Inches(3.05),midy, Inches(3.3),midy, width=2.4)                 # producer -> kafka
connector(s, Inches(5.15),midy, Inches(5.4),midy, width=2.4)                 # kafka -> consumer
connector(s, Inches(9.75),Inches(2.375), Inches(10.0),Inches(2.375), color=AMBER, width=2.2)   # consumer -> alerts
connector(s, Inches(9.75),Inches(3.4), Inches(10.0),Inches(3.4), color=DB, width=2.2)          # consumer -> audit
connector(s, Inches(9.75),Inches(4.15), Inches(10.0),Inches(4.28), color=RGBColor(0x41,0x58,0x6B), width=2.0)  # consumer -> log
connector(s, Inches(7.575),Inches(5.05), Inches(7.575),Inches(4.55), color=ACCENT, width=2.2, dashed=True, arrow=True)  # DB -> consumer (enrich)
textbox(s, Inches(8.05), Inches(4.62), Inches(1.6), Inches(0.3), "enrich lookup", size=10.5, color=ACCENT, bold=True)
# callout banner
cb = box(s, Inches(0.55), Inches(6.28), Inches(12.25), Inches(0.72), ACCENT, radius=0.10)
boxtext(s, cb, [("Flogo runs on BOTH sides of Kafka — producing AND consuming events.",16,WHITE,True)])
footer(s)

# =========================================================================
# 6 THE EVENT-FLOW PATTERN
s = slide(); eyebrow(s,"Reusable pattern"); title(s,"The event-flow pattern")
stages=[("1","Ingest",PRODUCER,"Consume events from Kafka / JMS at the edge, close to the producers."),
        ("2","Enrich",ACCENT,"Look up context from PostgreSQL or an API — setpoint, tolerance, cargo, customer."),
        ("3","Detect",ALERT,"Apply rules: temp above / below limits, or power not ON — flag breaches."),
        ("4","Route / Alert",AMBER,"Normal → log; breach → publish reefer.alerts, persist audit, notify.")]
cw=Inches(2.7); gap=Inches(0.3); x0=Inches(0.82); y0=Inches(2.25); h=Inches(3.0)
pos=[]
for i,(n,nm,cl,body) in enumerate(stages):
    x=x0+i*(cw+gap); pos.append(x)
    card(s,x,y0,cw,h,cl)
    bdg=s.shapes.add_shape(MSO_SHAPE.OVAL, _o(x+Inches(0.28)),_o(y0+Inches(0.3)),_e(Inches(0.55)),_e(Inches(0.55)))
    bdg.fill.solid(); bdg.fill.fore_color.rgb=cl; bdg.line.fill.background(); bdg.shadow.inherit=False
    boxtext(s,bdg,[(n,20,WHITE,True)])
    textbox(s,x+Inches(0.28),y0+Inches(1.05),cw-Inches(0.5),Inches(0.6),nm,size=20,color=INK,bold=True)
    textbox(s,x+Inches(0.28),y0+Inches(1.75),cw-Inches(0.5),Inches(1.1),body,size=14,color=MUTED,space_after=0)
    if i>0:
        connector(s, pos[i-1]+cw, y0+Inches(1.5), x, y0+Inches(1.5), color=RGBColor(0x8A,0x9A,0xAA), width=2.4)
textbox(s,Inches(0.8),Inches(5.65),Inches(11.7),Inches(1.0),
        [("One pattern, any event stream:  ",{"bold":True,"color":INK}),
         ("IoT / reefer telemetry, cargo & order events, gate / crane events, security events — "
          "the same Ingest → Enrich → Detect → Route flow is reused end to end.",{"color":MUTED})], size=16)
footer(s)

# =========================================================================
# 7 WHY FLOGO FOR EVENT HANDLING
s = slide(); eyebrow(s,"Why Flogo"); title(s,"Why Flogo for event handling")
feats=[("Tiny footprint",PRODUCER,"A single self-contained binary — minimal memory and CPU."),
       ("Low-code designer",CONSUMER,"Build and change event flows fast in a visual flow designer."),
       ("100+ connectors",ACCENT,"Kafka, JMS / EMS, databases, REST and cloud, out of the box."),
       ("Runs anywhere",DB,"Edge, container, Kubernetes or serverless — deploy close to events."),
       ("Event-native",AMBER,"Fast startup with triggers for Kafka, JMS, timer, HTTP and MQTT."),
       ("Config-driven",STREAM,"Topics, brokers and thresholds via app properties — retarget without code changes.")]
cw=Inches(3.8); gap=Inches(0.25); x0=Inches(0.8); rh=Inches(1.9); ry=[Inches(2.0),Inches(4.2)]
for i,(k,cl,body) in enumerate(feats):
    col=i%3; row=i//3
    x=x0+col*(cw+gap); y=ry[row]
    card(s,x,y,cw,rh,cl)
    chip(s,x+Inches(0.28),y+Inches(0.3),k,cl,w=Inches(2.4))
    textbox(s,x+Inches(0.28),y+Inches(0.95),cw-Inches(0.5),Inches(0.85),body,size=14.5,color=MUTED,space_after=0)
footer(s)

# =========================================================================
# 8 SUMMARY & NEXT STEPS (dark)
s = slide(NAVY)
band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0,0, _e(SW), _e(Inches(0.14)))
band.fill.solid(); band.fill.fore_color.rgb=ACCENT; band.line.fill.background(); band.shadow.inherit=False
textbox(s,Inches(0.9),Inches(0.95),Inches(11),Inches(0.4),"SUMMARY & NEXT STEPS",size=14,color=RGBColor(0x7F,0xD4,0xD1),bold=True)
textbox(s,Inches(0.9),Inches(1.4),Inches(11.6),Inches(1.0),"Complementary today, extensible tomorrow",size=34,color=WHITE,bold=True)
cw=Inches(5.9); h=Inches(3.7); y0=Inches(2.6)
# left card — what the demo proves
b=box(s,Inches(0.8),y0,cw,h,CARDDK,radius=0.06,line=NAVY2)
textbox(s,Inches(1.1),y0+Inches(0.28),cw-Inches(0.6),Inches(0.4),"WHAT THE REEFER DEMO PROVES",size=13,color=RGBColor(0x7F,0xD4,0xD1),bold=True)
bullets(s,Inches(1.1),y0+Inches(0.9),cw-Inches(0.6),Inches(2.6),
        ["Flogo produces AND consumes Kafka events",
         "Enriches telemetry from PostgreSQL (setpoint, tolerance, cargo)",
         "Detects breaches and routes alerts to reefer.alerts + audit table",
         "Complements TIBCO Streaming — it does not compete with it"],
        size=15.5, color=RGBColor(0xD7,0xE6,0xEF), dot=RGBColor(0x7F,0xD4,0xD1))
# right card — next steps
b=box(s,Inches(6.9),y0,cw,h,CARDDK,radius=0.06,line=NAVY2)
textbox(s,Inches(7.2),y0+Inches(0.28),cw-Inches(0.6),Inches(0.4),"NEXT STEPS FOR PSA",size=13,color=RGBColor(0x7F,0xD4,0xD1),bold=True)
bullets(s,Inches(7.2),y0+Inches(0.9),cw-Inches(0.6),Inches(2.6),
        ["Point the agents at PSA's real Kafka / EMS brokers",
         "Add a JMS / EMS bridge flow",
         "Wire Flogo into the existing TIBCO Streaming pipeline (pre-process & act-on)",
         "Deploy agents to the edge / Kubernetes",
         "Extend cargo profiles and alert rules"],
        size=15.5, color=RGBColor(0xD7,0xE6,0xEF), dot=ACCENT)
textbox(s,Inches(0.9),Inches(6.5),Inches(11.5),Inches(0.4),
        "Flogo_Apps/Reefer_ColdChain_Monitoring_Use_Case",size=13,color=RGBColor(0x9F,0xD8,0xD5))
footer(s,"Thank you",dark=True)

# =========================================================================
out_dir = os.path.dirname(os.path.abspath(__file__))
os.makedirs(out_dir, exist_ok=True)
out = os.path.join(out_dir, "ReeferColdChain_Architecture.pptx")
prs.save(out)
print("saved", out, "slides:", len(prs.slides._sldIdLst))
